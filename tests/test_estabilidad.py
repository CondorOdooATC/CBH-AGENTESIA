"""Casos de aceptación de la versión estable (revisión v1.2 → v1.3): coherencia de cálculos, propuestas,
aprobaciones y operaciones de principio a fin."""
from __future__ import annotations

import numpy as np
import pandas as pd
import pytest

from app import db
from app.agents import autonomia, planificador
from app.agents import demanda as agente_demanda
from app.ml import pronostico as P
from app.odoo import acciones as OA
from app.odoo.client import OdooError

HOY = pd.Timestamp("2026-09-12")


def _limpiar_pendientes(agente=None):
    """Las pruebas comparten base: se rechazan las propuestas pendientes previas para medir sólo lo de la prueba."""
    for e in autonomia.ESTADOS_PENDIENTES:
        for a in db.acciones(estado=e, agente=agente, limite=5000):
            db.transicion_accion(a["id"], e, "rechazada", resultado="limpieza de prueba")


def _consumo_sintetico(dias=120, por_dia=10.0, folios_dia=2, hospital="HGZ 17 Monterrey", loc="HGZ17/Stock", pid=1001,
                       producto="Sevoflurano 250 mL frasco", unidad="mL"):
    filas = []
    fecha0 = HOY - pd.Timedelta(days=dias)
    n = 0
    for d in range(dias):
        f = fecha0 + pd.Timedelta(days=d)
        if f.weekday() >= 5:
            continue
        for k in range(folios_dia):
            n += 1
            filas.append({"fecha": f + pd.Timedelta(hours=9 + k), "folio": f"F{n:05d}", "hospital": hospital, "almacen": loc, "subalmacen": loc,
                          "producto_id": pid, "producto": producto, "unidad": unidad, "cantidad": por_dia / folios_dia, "importe": por_dia / folios_dia * 15.0,
                          "medico": "Dr. A", "auxiliar": "Aux B"})
    return pd.DataFrame(filas)


def _existencias(stock=200.0, loc="HGZ17/Stock", pid=1001, producto="Sevoflurano 250 mL frasco", extra=None):
    filas = [{"producto_id": pid, "producto": producto, "ubicacion": loc, "almacen": loc, "cantidad": stock, "disponible": stock,
              "reservado": 0.0, "caducado": 0.0, "utilizable": stock, "lote": "", "caducidad": None}]
    for e in (extra or []):
        filas.append({"producto_id": pid, "producto": producto, "ubicacion": e[0], "almacen": e[0], "cantidad": e[1], "disponible": e[1],
                      "reservado": 0.0, "caducado": 0.0, "utilizable": e[1], "lote": "", "caducidad": None})
    return pd.DataFrame(filas)


def _agenda(dia: int, folios: int, hospital="HGZ 17 Monterrey"):
    return pd.DataFrame([{"hospital": hospital, "hospital_id": 102, "dia": (HOY + pd.Timedelta(days=dia)).date(), "folios": folios}])


def _cfg(**kw):
    c = P.ConfigPronostico(horizonte_dias=30, historia_dias=120, hoy=HOY, umbral_sin_captura=0.0, es_cedis={"CEDIS-MTY/Stock"})
    for k, v in kw.items():
        setattr(c, k, v)
    return c


# ── 1 · Las fechas reales de los folios programados se respetan ─────────────
def test_1_agenda_en_su_fecha_real():
    df = _consumo_sintetico()
    ex = _existencias(stock=200.0)
    r1 = P.pronosticar(df, ex, _cfg(), folios_prog=_agenda(1, 20))
    r20 = P.pronosticar(df, ex, _cfg(), folios_prog=_agenda(20, 20))
    p1 = r1["pronosticos"][r1["pronosticos"]["almacen"] == "HGZ17/Stock"].sort_values("fecha")
    p20 = r20["pronosticos"][r20["pronosticos"]["almacen"] == "HGZ17/Stock"].sort_values("fecha")
    # la demanda de la agenda cae exactamente en el día programado (índice 0 = mañana)
    assert p1["demanda_agenda"].iloc[0] > 0 and p1["demanda_agenda"].iloc[19] == 0
    assert p20["demanda_agenda"].iloc[19] > 0 and p20["demanda_agenda"].iloc[0] == 0
    assert abs(p1["demanda_agenda"].sum() - p20["demanda_agenda"].sum()) < 1e-6      # misma necesidad, otra fecha
    # mañana la demanda proyectada es distinta según dónde esté la jornada
    assert p1["demanda_proyectada"].iloc[0] > p20["demanda_proyectada"].iloc[0]
    l1 = r1["resurtido"][(r1["resurtido"]["nivel"] == "local")].iloc[0]
    l20 = r20["resurtido"][(r20["resurtido"]["nivel"] == "local")].iloc[0]
    # el quiebre se desplaza con la agenda y la fecha necesaria de abastecimiento es anterior al quiebre
    assert l1["dia_quiebre"] is not None and l20["dia_quiebre"] is not None and l1["dia_quiebre"] < l20["dia_quiebre"]
    assert l1["dia_necesario"] < l1["dia_quiebre"] and l1["fecha_necesaria"] < l1["fecha_quiebre"]
    # la agenda es un piso: no se suma a la demanda base del día
    base = p1["pronostico"].iloc[0]
    assert abs(p1["demanda_proyectada"].iloc[0] - max(base, p1["demanda_agenda"].iloc[0])) < 1e-6
    # cancelar la agenda vuelve al pronóstico base sin residuos
    r0 = P.pronosticar(df, ex, _cfg(), folios_prog=None)
    p0 = r0["pronosticos"][r0["pronosticos"]["almacen"] == "HGZ17/Stock"].sort_values("fecha")
    assert (p0["demanda_agenda"] == 0).all() and abs(p0["demanda_proyectada"].sum() - p0["pronostico"].sum()) < 1e-6


# ── 2 · La jornada extraordinaria no se cuenta dos veces ────────────────────
def test_2_agenda_una_sola_vez_y_sin_duplicar_propuestas(sim):
    df = _consumo_sintetico()
    ex = _existencias(stock=60.0, extra=[("CEDIS-MTY/Stock", 5000.0)])
    cfg = _cfg()
    res = P.pronosticar(df, ex, cfg, folios_prog=_agenda(3, 20), compromisos=autonomia.compromisos_activos())
    loc = res["resurtido"][(res["resurtido"]["nivel"] == "local") & (res["resurtido"]["almacen"] == "HGZ17/Stock")].iloc[0]
    pr = res["pronosticos"][res["pronosticos"]["almacen"] == "HGZ17/Stock"]
    assert abs(loc["ajuste_agenda"] - float(np.maximum(0, pr["demanda_agenda"] - pr["pronostico"]).sum())) < 1e-6
    assert loc["demanda_horizonte"] == pytest.approx(float(pr["demanda_proyectada"].sum()), rel=1e-3)
    # el planificador determinista NO vuelve a multiplicar la demanda: sus escenarios son hipotéticos e incluyen las propuestas
    cid = db.iniciar_corrida("demanda", {})
    ctx = planificador.ContextoPlan(res, df, pd.DataFrame(columns=["tipo", "ref", "retrasada", "confirmada", "producto", "cantidad", "cantidad_doc", "unidad_doc", "fecha_prevista"]),
                                    _agenda(3, 20), [], cid, cfg)
    esc_base = ctx.simular_escenario("sin cambios")
    assert esc_base["empeoran"] == 0                                    # el mismo plan no "empeora" nada
    esc = ctx.simular_escenario("agenda +20 % sólo el día 3", {"HGZ 17": 1.2}, dias_desde=3, dias_hasta=3)
    esc_fuera = ctx.simular_escenario("+20 % fuera de la agenda", {"HGZ 17": 1.2}, dias_desde=25, dias_hasta=25)
    for d_ in esc["detalle"]:
        assert d_["faltante_residual"] <= d_["faltante_estimado"]
    assert (esc_fuera["detalle"][0]["faltante_residual"] if esc_fuera["detalle"] else 0) <= (esc["detalle"][0]["faltante_residual"] if esc["detalle"] else 0)
    plan = planificador.plan_determinista(ctx)
    assert not any("×" in (d_.get("decision") or "") for d_ in plan["decisiones"])
    assert any("ya lo incorpora" in n for n in plan["lo_que_el_modelo_no_ve"])
    # repetir el análisis sin cambios no infla cantidades ni duplica propuestas
    _limpiar_pendientes()
    res = P.pronosticar(df, ex, cfg, folios_prog=_agenda(3, 20), compromisos=autonomia.compromisos_activos())
    acc1 = agente_demanda._proponer(res, cid, "t")
    cid2 = db.iniciar_corrida("demanda", {})
    res2 = P.pronosticar(df, ex, cfg, folios_prog=_agenda(3, 20), compromisos=autonomia.compromisos_activos())
    acc2 = agente_demanda._proponer(res2, cid2, "t")
    n2 = len([a for e in autonomia.ESTADOS_PENDIENTES for a in db.acciones(estado=e, agente="demanda")])
    assert len(acc1) >= 1 and len(acc2) == len(acc1) and all(a.get("reutilizada") for a in acc2)
    assert n2 == len(acc1)
    ids = {a["id"] for a in acc1}
    assert {a["id"] for a in acc2} == ids
    for a in acc2:
        assert db.accion(a["id"])["version"] == 1                       # sin cambio material → sin versión nueva


# ── 3 · Existencias compartidas entre todas las propuestas ──────────────────
def test_3_asignacion_conjunta_no_compromete_mas_de_lo_que_hay():
    unidad = "pz"
    filas = []
    for alm, crit, sug, q in (("HGZ17/Stock", "critico", 300.0, 2), ("HGZ67/Stock", "critico", 250.0, 3)):
        filas.append({"producto_id": 1010, "producto": "Jeringa 10 mL", "almacen": alm, "hospital": alm, "unidad": unidad, "criticidad": crit,
                      "sugerido": sug, "dia_quiebre": q, "fecha_quiebre": "2026-09-14", "dia_necesario": q - 1, "fecha_necesaria": "2026-09-13",
                      "stock_actual": 5.0, "stock_proyectado": 5.0, "demanda_diaria": 20.0, "demanda_horizonte": 600.0, "punto_reorden": 60.0,
                      "dias_cobertura": 0.3, "es_cedis": False, "costo_unit": 4.0, "nivel": "local", "stock_seguridad": 20.0})
    filas.append({"producto_id": 1010, "producto": "Jeringa 10 mL", "almacen": "CEDIS-MTY/Stock", "hospital": "", "unidad": unidad, "criticidad": "fuente",
                  "sugerido": 0.0, "dia_quiebre": None, "fecha_quiebre": None, "dia_necesario": None, "fecha_necesaria": None,
                  "stock_actual": 400.0, "stock_proyectado": 400.0, "demanda_diaria": 0.0, "demanda_horizonte": 0.0, "punto_reorden": 0.0,
                  "dias_cobertura": None, "es_cedis": True, "costo_unit": 4.0, "nivel": "local", "stock_seguridad": 0.0})
    local = pd.DataFrame(filas)
    cfg = _cfg(reserva_cedis_dias=0)
    reb = P.proponer_rebalanceo(local, {}, compromisos=None, cfg=cfg)
    assert reb["cantidad"].sum() <= 400 and reb["cantidad"].sum() == 400      # 300 + 100, nunca 550
    assert reb.iloc[0]["destino"] == "HGZ17/Stock" and reb.iloc[0]["cantidad"] == 300   # el más urgente primero
    # con reserva operativa del CEDIS (7 días × 40/día = 280) "sin consumo propio" no significa "todo libre"
    reb2 = P.proponer_rebalanceo(local, {}, compromisos=None, cfg=_cfg(reserva_cedis_dias=7))
    assert reb2["cantidad"].sum() <= 400 - 280 and (reb2["reserva_origen"] == 280).all()
    # lo que otras propuestas activas ya apartaron también se descuenta y lo ya cubierto en el destino no se vuelve a pedir
    reb3 = P.proponer_rebalanceo(local, {}, compromisos={"origen": {(1010, "CEDIS-MTY/Stock"): 150.0}, "destino": {(1010, "HGZ17/Stock"): 100.0}}, cfg=cfg)
    assert reb3["cantidad"].sum() <= 250 and reb3[reb3["destino"] == "HGZ17/Stock"]["cantidad"].sum() <= 200
    # cantidades enteras para piezas
    assert all(float(c).is_integer() for c in reb3["cantidad"])


def test_3b_aprobaciones_en_serie_respetan_el_libro(sim):
    """Dos transferencias del mismo origen aprobadas en grupo: la segunda no puede llevarse lo que ya se llevó la primera."""
    origen = next(l for l in sim.tablas["stock.location"] if l["complete_name"] == "CEDIS-LR/Stock")
    dest = next(l for l in sim.tablas["stock.location"] if l["complete_name"] == "HGZ67/Stock")
    pid = 1004
    disponible = sum(float(q["quantity"]) - float(q.get("reserved_quantity") or 0) for q in sim.tablas["stock.quant"]
                     if q["product_id"][0] == pid and q["location_id"][0] == origen["id"])
    assert disponible > 10
    base = {"producto_id": pid, "producto": "Propofol", "unidad": "pz", "origen_id": origen["id"], "destino_id": dest["id"],
            "origen": "CEDIS-LR/Stock", "destino": "HGZ67/Stock"}
    a = autonomia.proponer("demanda", "transferencia_interna", "", {**base, "cantidad": round(disponible * 0.7)}, impacto={"importe": 10}, sincronizar=False)
    b = autonomia.proponer("demanda", "transferencia_interna", "", {**base, "cantidad": round(disponible * 0.7)}, impacto={"importe": 10}, sincronizar=False)
    res = autonomia.aprobar_varias([a["id"], b["id"]], "ana", "admin")
    assert res[0]["estado"] == "ejecutada"
    assert res[1]["estado"] in ("requiere_revision", "propuesta")            # se topó a lo que quedaba; requiere nueva aprobación
    b2 = db.accion(b["id"])
    assert b2["payload"]["cantidad"] <= disponible - round(disponible * 0.7) + 1e-6 and "comprometidos" in (b2["revalidacion"] or "")


# ── 5 · Unidades y redondeo coherentes de punta a punta ─────────────────────
def test_5_unidad_y_redondeo_titulo_payload_odoo(sim):
    origen = next(l for l in sim.tablas["stock.location"] if l["complete_name"] == "CEDIS-LR/Stock")
    dest = next(l for l in sim.tablas["stock.location"] if l["complete_name"] == "HGZ17/Stock")
    r = autonomia.proponer("demanda", "transferencia_interna", "", {"producto_id": 1004, "producto": "Propofol 200 mg/20 mL ampolleta", "unidad": "pz",
                                                                    "cantidad": 40.7, "origen_id": origen["id"], "destino_id": dest["id"],
                                                                    "origen": "CEDIS-LR/Stock", "destino": "HGZ17/Stock"},
                           impacto={"importe": 100, "costo_unit": 2.0}, sincronizar=False)
    a = db.accion(r["id"])
    assert a["payload"]["cantidad"] == 41 and "41 pz" in a["titulo"] and a["impacto"]["cantidad"] == 41 and a["impacto"]["importe"] == 82.0
    assert "41 pz" in a["efecto"] and " u)" not in a["efecto"]
    out = autonomia.aprobar(r["id"], "ana", "admin")
    assert out["estado"] == "ejecutada"
    move = next(m for m in sim.tablas["stock.move"] if m.get("picking_id") and m["picking_id"][0] == out["odoo"]["id"])
    assert move["product_uom_qty"] == 41.0                                   # lo enviado a Odoo es lo aprobado
    assert db.accion(r["id"])["cantidad_aprobada"] == 41 and db.accion(r["id"])["cantidad_ejecutada"] == 41
    # mililitros conservan un decimal; la RFQ se pide en frascos enteros
    assert P.redondear(1157.26, "mL", None, arriba=True) == 1157.3 and P.redondear(402.7, "pz", None, arriba=True) == 403
    rfq = OA.crear_solicitud_compra(1001, 7321, referencia="Agente IA · prueba-uom", cli=sim)
    assert rfq["cantidad_doc"] == 30 and rfq["unidad_doc"] == "Frasco 250 mL"
    # sin conversión conocida, la propuesta de compra se bloquea en vez de suponer equivalencias
    b = autonomia.proponer("demanda", "solicitud_compra", "", {"producto_id": 1001, "producto": "Sevoflurano", "unidad": "mL", "cantidad": 500,
                                                               "unidad_compra": "Frasco 250 mL", "conversion_faltante": True}, sincronizar=False)
    assert b["estado"] == "bloqueada" and "conversión" in " ".join(b["motivos"])


# ── 6 · El impacto se recalcula después de cada ajuste ──────────────────────
def test_6_ajuste_recalcula_impacto(sim):
    df = _consumo_sintetico()
    ex = _existencias(stock=60.0, extra=[("CEDIS-MTY/Stock", 5000.0)])
    cfg = _cfg()
    res = P.pronosticar(df, ex, cfg, folios_prog=_agenda(3, 20))
    cid = db.iniciar_corrida("demanda", {})
    origen = next(l for l in sim.tablas["stock.location"] if l["complete_name"] == "CEDIS-MTY/Stock")
    dest = next(l for l in sim.tablas["stock.location"] if l["complete_name"] == "HGZ17/Stock")
    r = autonomia.proponer("demanda", "transferencia_interna", "", {"producto_id": 1001, "producto": "Sevoflurano 250 mL frasco", "unidad": "mL", "cantidad": 100,
                                                                    "origen_id": origen["id"], "destino_id": dest["id"], "origen": "CEDIS-MTY/Stock", "destino": "HGZ17/Stock"},
                           impacto={"importe": 1500, "costo_unit": 15.0, "stock_origen": 5000}, corrida_id=cid, sincronizar=False)
    ctx = planificador.ContextoPlan(res, df, pd.DataFrame(), _agenda(3, 20), [], cid, cfg)
    antes = ctx._proyectar_local(1001, "HGZ17/Stock", extra_in=[(2, 100)])
    aj = ctx.ajustar_propuesta(r["id"], 400, "escenario")
    a = db.accion(r["id"])
    assert aj["ok"] and a["version"] == 2 and a["impacto"]["importe"] == 6000.0 and a["impacto"]["cantidad"] == 400 and "400 mL" in a["titulo"]
    despues = ctx._proyectar_local(1001, "HGZ17/Stock", extra_in=[(2, 400)])
    assert a["impacto"]["cobertura_destino_despues"] == despues["cobertura_dias"] >= antes["cobertura_dias"]
    assert a["impacto"]["faltante_residual_destino"] == round(despues["faltante_para_seguridad"], 1)
    assert a["impacto"]["disponible_origen"] is not None and a["impacto"]["reserva_origen"] is not None


# ── 7 · Seguridad de aprobaciones y ejecuciones ─────────────────────────────
def test_7_revalida_antes_de_ejecutar_y_detiene_si_no_puede_consultar(sim, monkeypatch):
    origen = next(l for l in sim.tablas["stock.location"] if l["complete_name"] == "CEDIS-LR/Stock")
    dest = next(l for l in sim.tablas["stock.location"] if l["complete_name"] == "HGZ17/Stock")
    payload = {"producto_id": 1004, "producto": "Propofol", "unidad": "pz", "cantidad": 3, "origen_id": origen["id"], "destino_id": dest["id"],
               "origen": "CEDIS-LR/Stock", "destino": "HGZ17/Stock"}
    r = autonomia.proponer("demanda", "transferencia_interna", "", payload, impacto={"importe": 10}, sincronizar=False)
    n0 = len(sim.tablas["stock.picking"])
    real = sim.search_read

    def roto(modelo, *a, **k):
        if modelo == "stock.quant":
            raise OdooError("Odoo no responde")
        return real(modelo, *a, **k)
    monkeypatch.setattr(sim, "search_read", roto)
    out = autonomia.aprobar(r["id"], "ana", "admin")
    assert out["estado"] == "requiere_revision" and len(sim.tablas["stock.picking"]) == n0     # nada se escribió
    monkeypatch.setattr(sim, "search_read", real)
    # un cambio de destino tras la aprobación invalida la aprobación (versión y efecto exacto)
    r2 = autonomia.proponer("demanda", "transferencia_interna", "", {**payload, "cantidad": 4}, impacto={"importe": 10}, sincronizar=False)
    db.actualizar_accion(r2["id"], estado="aprobada", version_aprobada=1, version=1, efecto_aprobado="otro efecto")
    res = autonomia.ejecutar(r2["id"])
    assert res["estado"] == "propuesta" and db.accion(r2["id"])["aprobado_por"] is None


def test_7b_produccion_no_ejecuta_acciones_no_validadas(sim, monkeypatch):
    from app.config import settings
    monkeypatch.setattr(settings, "APP_ENV", "production")
    r = autonomia.proponer("vigilancia", "confirmar_transferencia", "Confirmar CEDIS/INT/00912", {"picking": "CEDIS/INT/00912"}, sincronizar=False)
    assert "ejecución manual" in db.accion(r["id"])["efecto"].lower() or "manual" in db.accion(r["id"])["efecto"].lower()
    out = autonomia.aprobar(r["id"], "ana", "admin")
    assert out["estado"] == "aprobada_manual"
    pk = next(p for p in sim.tablas["stock.picking"] if p["name"] == "CEDIS/INT/00912")
    assert pk["state"] != "done"


# ── 8 · Las acciones nuevas se validan una por una ──────────────────────────
def test_8_cuarentena_bloquea_de_verdad(sim):
    q = next(x for x in sim.tablas["stock.quant"] if x.get("lot_id") and float(x["quantity"]) >= 5 and float(x.get("reserved_quantity") or 0) == 0)
    pid, loc, lote = q["product_id"][0], q["location_id"][0], q["lot_id"][1]
    res = OA.cuarentena_lote(pid, lote, 5, loc, referencia="Agente IA · cuarentena-test", cli=sim)
    assert res["bloqueado"] is True and res["estado"] == "assigned" and res["reservado"] == 5
    assert float(q["reserved_quantity"]) == 5.0                                # la reserva es el bloqueo real
    # pedir más de lo que hay NO se reporta como bloqueado
    res2 = OA.cuarentena_lote(pid, lote, float(q["quantity"]) + 100, loc, referencia="Agente IA · cuarentena-test-2", cli=sim)
    assert res2["bloqueado"] is False and res2["advertencia"]
    # un lote inexistente en un producto con lotes se rechaza
    with pytest.raises(OdooError):
        OA.cuarentena_lote(pid, "LOTE-QUE-NO-EXISTE", 1, loc, referencia="Agente IA · cuarentena-test-3", cli=sim)


def test_8b_validar_recepcion_verifica_y_no_supone(sim):
    origen = next(l for l in sim.tablas["stock.location"] if l["complete_name"] == "CEDIS-LR/Stock")
    dest = next(l for l in sim.tablas["stock.location"] if l["complete_name"] == "HGZ17/Stock")
    # borrador → no se valida (requiere decisión / preparación)
    t = OA.crear_transferencia_interna(1004, 2, origen["id"], dest["id"], referencia="Agente IA · val-1", cli=sim)
    with pytest.raises(OA.RequiereDecision):
        OA.validar_recepcion(t["ref"], cli=sim)
    # confirmada y reservada con cantidades completas → done y las existencias se mueven
    antes_o = sum(float(x["quantity"]) for x in sim.tablas["stock.quant"] if x["product_id"][0] == 1004 and x["location_id"][0] == origen["id"])
    antes_d = sum(float(x["quantity"]) for x in sim.tablas["stock.quant"] if x["product_id"][0] == 1004 and x["location_id"][0] == dest["id"])
    c = OA.confirmar_transferencia(t["ref"], cli=sim)
    assert c["estado"] == "assigned"
    v = OA.validar_recepcion(t["ref"], cli=sim)
    assert v["estado"] == "done"
    despues_o = sum(float(x["quantity"]) for x in sim.tablas["stock.quant"] if x["product_id"][0] == 1004 and x["location_id"][0] == origen["id"])
    despues_d = sum(float(x["quantity"]) for x in sim.tablas["stock.quant"] if x["product_id"][0] == 1004 and x["location_id"][0] == dest["id"])
    assert despues_o == pytest.approx(antes_o - 2) and despues_d == pytest.approx(antes_d + 2)
    # entrega parcial: Odoo devuelve un asistente → decisión humana, sin marcar éxito
    t2 = OA.crear_transferencia_interna(1004, 5, origen["id"], dest["id"], referencia="Agente IA · val-2", cli=sim)
    OA.confirmar_transferencia(t2["ref"], cli=sim)
    mv = next(m for m in sim.tablas["stock.move"] if m.get("picking_id") and m["picking_id"][0] == t2["id"])
    mv["quantity"] = 3.0
    with pytest.raises(OA.RequiereDecision):
        OA.validar_recepcion(t2["ref"], cli=sim)
    pk = next(p for p in sim.tablas["stock.picking"] if p["id"] == t2["id"])
    assert pk["state"] != "done"
    # a través de la autonomía queda en requiere_revision (no en error ni en ejecutada)
    r = autonomia.proponer("vigilancia", "validar_recepcion", "Validar", {"picking": t2["ref"]}, sincronizar=False)
    autonomia.set_politicas({"doble_aprobacion_riesgo_alto": False})
    out = autonomia.aprobar(r["id"], "ana", "admin")
    autonomia.set_politicas({"doble_aprobacion_riesgo_alto": True})
    assert out["estado"] == "requiere_revision" and "decisión" in out["mensaje"].lower() or "parcial" in out["mensaje"].lower()


def test_8c_fallo_parcial_liga_el_documento_sin_duplicar(sim, monkeypatch):
    # producto con existencia holgada en el origen y sin compromisos firmes de otras pruebas (1015: 354 pz en CEDIS-LR)
    _limpiar_pendientes()
    origen = next(l for l in sim.tablas["stock.location"] if l["complete_name"] == "CEDIS-LR/Stock")
    dest = next(l for l in sim.tablas["stock.location"] if l["complete_name"] == "HGZ17/Stock")
    r = autonomia.proponer("demanda", "transferencia_interna", "", {"producto_id": 1015, "producto": "Cánula de Guedel #4", "unidad": "pz", "cantidad": 2,
                                                                    "origen_id": origen["id"], "destino_id": dest["id"], "origen": "CEDIS-LR/Stock", "destino": "HGZ17/Stock"},
                           impacto={"importe": 10}, sincronizar=False)
    real = sim.mensaje_chatter

    def falla(*a, **k):
        raise OdooError("caída después de crear")
    monkeypatch.setattr(sim, "mensaje_chatter", falla)
    monkeypatch.setattr(OA, "crear_transferencia_interna", lambda *a, **k: (_ for _ in ()).throw(OdooError("caída después de crear")) if not sim.search_read("stock.picking", [["origin", "=", k.get("referencia")]], ["id"], limite=1)
                        else (_ for _ in ()).throw(OdooError("caída después de crear")))
    # simulamos que Odoo sí creó el documento antes de la caída
    sim.create("stock.picking", {"name": "CEDIS/INT/TEST-PARCIAL", "origin": f"Agente IA · acción #{r['id']}", "state": "draft",
                                 "location_id": origen["id"], "location_dest_id": dest["id"]})
    n0 = len(sim.tablas["stock.picking"])
    out = autonomia.aprobar(r["id"], "ana", "admin")
    monkeypatch.setattr(sim, "mensaje_chatter", real)
    a = db.accion(r["id"])
    assert out["estado"] == "ejecutada" and a["odoo_ref"] == "CEDIS/INT/TEST-PARCIAL" and "falló" in (a["revalidacion"] or "")
    assert len(sim.tablas["stock.picking"]) == n0                              # no se duplicó


# ── 10 · Regresiones adicionales ────────────────────────────────────────────
def test_10_excel_recortado_lo_dice(sim, monkeypatch, tmp_path):
    from app.reports import builders
    from app.reports.excel import LibroExcel
    from openpyxl import load_workbook
    monkeypatch.setattr(LibroExcel, "MAX_FILAS", 50)
    df = pd.DataFrame({"producto": [f"P{i}" for i in range(120)], "cantidad": range(120), "unidad": ["pz"] * 120})
    r = builders.reporte_dataframe("Prueba de recorte", df, tipo="consulta", hoja="Datos", usuario="t")
    assert r["recortado"] is True and r["recortes"][0]["mostradas"] == 50 and r["recortes"][0]["totales"] == 120
    wb = load_workbook(r["ruta"])
    acerca = {row[0].value: row[1].value for row in wb["Acerca de"].iter_rows(min_row=3) if row[0].value}
    assert "NO" in str(acerca.get("Completo")) and "50" in str(acerca.get("Completo"))
    hoja = wb["Datos"]
    assert "RECORTADO" in str(hoja["A2"].value)
    # un reporte completo lo dice también
    r2 = builders.reporte_dataframe("Prueba completa", df.head(10), tipo="consulta", hoja="Datos", usuario="t")
    assert r2["recortado"] is False
    wb2 = load_workbook(r2["ruta"])
    acerca2 = {row[0].value: row[1].value for row in wb2["Acerca de"].iter_rows(min_row=3) if row[0].value}
    assert str(acerca2.get("Completo")).startswith("sí")


def test_10b_privacidad_estricta_de_conversaciones(sim):
    from fastapi.testclient import TestClient
    from app.main import app
    db.crear_usuario("priv1", "test1234", "Priv 1", "consulta")
    db.crear_usuario("adm_priv", "test1234", "Admin", "admin")
    with TestClient(app) as c:
        c.post("/login", data={"usuario": "priv1", "password": "test1234", "next": "/"}, follow_redirects=False)
        cid = c.post("/api/conversaciones").json()["id"]
    with TestClient(app) as c2:
        c2.post("/login", data={"usuario": "adm_priv", "password": "test1234", "next": "/"}, follow_redirects=False)
        assert c2.get(f"/copiloto?c={cid}").status_code == 403          # ni siquiera un administrador
        assert c2.post("/api/chat", json={"conversacion_id": cid, "texto": "hola"}).status_code == 403


def test_10c_aclaracion_no_es_hecho_ni_regla(sim):
    """Una explicación aportada en el chat de un caso queda como declaración trazable; no cambia hechos ni crea reglas."""
    from app.agents import copiloto
    c = db.casos(limite=1)
    if not c:
        pytest.skip("sin casos")
    caso = c[0]
    n_apr = len(db.aprendizaje(limite=1000))
    conv = db.nueva_conversacion("ana", "Caso prueba")
    r = copiloto.responder(conv, "El médico confirma que fue una cirugía de 5 horas con reintervención", "ana", "operacion", contexto={"caso_id": caso["id"]})
    assert "aclaración" in r["texto"].lower() and "no la convierto en hecho" in r["texto"].lower()
    acl = db.aclaraciones(caso["id"])
    assert acl and acl[0]["usuario"] == "ana" and acl[0]["alcance"] == "caso" and acl[0]["verificada"] == 0
    assert len(db.aprendizaje(limite=1000)) == n_apr                 # sin regla global
    assert db.caso(caso["id"])["expediente"] == caso["expediente"]  # los hechos no cambian
    # el expediente determinista la cita como declaración, aparte de los hechos
    from app.agents import investigador
    from app.odoo import queries
    df = queries.consumo(dias=180)
    ctx = investigador.Contexto(df)
    o = dict(caso.get("entidades") or {}); o.update({"tipo": "anomalia", "titulo": "t", "regla": "R01_EXCEDE_ENVASE", "severidad": "alta"})
    e = investigador.expediente_determinista(o, ctx)
    assert all("None" not in x and "nan" not in x for x in e["evidencia"])
    if e.get("aclaraciones"):
        assert e["aclaraciones"][0]["verificada"] is False
