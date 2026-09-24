"""Reportes Excel concretos de la plataforma (todos registrados en la bitácora)."""
from __future__ import annotations

import json
import re
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd

from .. import db
from ..config import settings
from .excel import (LibroExcel, SEMAFORO_ACCION, SEMAFORO_CRITICIDAD, SEMAFORO_ESTADO_CAD,
                    SEMAFORO_SEVERIDAD)

MXN = '"$"#,##0.00'
ENT = "#,##0"
DEC = "#,##0.00"
PCT = "0.0%"


def _slug(s: str) -> str:
    s = re.sub(r"[^A-Za-z0-9]+", "-", s).strip("-").lower()
    return s[:60] or "reporte"


def _ruta(tipo: str, titulo: str) -> Path:
    ts = datetime.now().strftime("%Y%m%d-%H%M%S")
    return settings.REPORTS_DIR / f"{ts}_{_slug(tipo)}_{_slug(titulo)}.xlsx"


def _registrar(libro: LibroExcel, ruta: Path, tipo: str, titulo: str, parametros: dict, filas: int,
               usuario: str = "", metadatos: dict | None = None) -> dict:
    from ..odoo import schema
    meta = {"Tipo de reporte": tipo, "Registros": f"{filas:,}", "Origen de datos": schema.modelo("consumo") or "Odoo",
            "Parámetros": parametros or "—"}
    meta.update(metadatos or {})
    libro.acerca_de(meta)          # "Completo: sí / NO — recortado…" lo calcula el libro con lo realmente escrito
    libro.guardar(ruta)
    recortado = bool(libro.recortes)
    rid = db.registrar_reporte(ruta.name, str(ruta), tipo, titulo, {**(parametros or {}), "recortado": recortado} if isinstance(parametros, dict) else parametros,
                               filas, ruta.stat().st_size, usuario)
    db.log("warn" if recortado else "info", "reportes", f"Excel generado: {ruta.name}" + (" (RECORTADO al límite de Excel)" if recortado else ""),
           f"tipo={tipo} filas={filas}", usuario)
    return {"id": rid, "archivo": ruta.name, "ruta": str(ruta), "tipo": tipo, "titulo": titulo,
            "filas": filas, "url": f"/reportes/{rid}/descargar", "recortado": recortado, "recortes": libro.recortes}


# ── Agente 1 ────────────────────────────────────────────────────────────────
def reporte_anomalias(resultado: dict, informe: str = "", corrida: dict | None = None,
                      usuario: str = "") -> dict:
    h: pd.DataFrame = resultado.get("hallazgos", pd.DataFrame())
    ag: pd.DataFrame = resultado.get("agregados", pd.DataFrame())
    riesgos = resultado.get("riesgos", {})
    hist = resultado.get("historico", {})
    bas = resultado.get("basculas", {})
    titulo = "Control Inteligente de Consumo"
    libro = LibroExcel(titulo, f"Agente 1 · Corrida #{(corrida or {}).get('id', '')} · "
                               f"{resultado.get('n_lineas', 0):,} líneas analizadas")
    sev = h["severidad"].value_counts().to_dict() if not h.empty else {}
    libro.portada(
        kpis=[("Líneas de consumo analizadas", int(resultado.get("n_lineas", 0))),
              ("Hallazgos por línea", int(len(h))),
              ("  · Críticos", int(sev.get("critica", 0))), ("  · Altos", int(sev.get("alta", 0))),
              ("  · Medios", int(sev.get("media", 0))),
              ("Patrones agregados", int(len(ag))),
              ("Importe en riesgo (MXN)", float(h["importe_riesgo"].sum()) if not h.empty else 0.0),
              ("Importe patrones (MXN)", float(ag["importe_riesgo"].sum()) if not ag.empty else 0.0),
              ("Líneas con báscula", int(bas.get("lineas_con_bascula", 0))),
              ("Lecturas imposibles de báscula", int(bas.get("imposibles", 0))),
              ("Discrepancias de báscula", int(bas.get("discrepancias", 0))),
              ("Gramos no explicados por báscula", float(bas.get("gramos_no_explicados", 0.0)))],
        notas=["Los hallazgos describen patrones estadísticos y de reglas; requieren verificación operativa "
               "antes de cualquier conclusión sobre personas.",
               "Severidad: crítica = actuar hoy; alta = revisar esta semana; media = monitorear.",
               "El importe en riesgo se calcula como (cantidad − esperado) × costo unitario."])
    if informe:
        libro.hoja_texto("Informe", informe, "Informe del Agente 1")
    if not h.empty:
        hh = h.copy()
        hh["metodos"] = hh["metodos"].map(lambda m: ", ".join(m))
        hh["motivos"] = hh["motivos"].map(lambda m: " | ".join(m))
        libro.hoja_tabla("Hallazgos", hh, columnas=[
            "severidad", "score", "fecha", "folio", "hospital", "subalmacen", "medico", "auxiliar", "producto",
            "lote", "cantidad", "unidad", "esperado", "desviacion", "importe_riesgo", "duracion_min", "tasa",
            "peso_inicial", "peso_final", "motivos", "metodos", "huella"],
            etiquetas={"score": "Puntaje", "subalmacen": "Sub-almacén", "medico": "Médico",
                       "duracion_min": "Duración (min)", "tasa": "Tasa mL/min", "importe_riesgo": "Importe en riesgo",
                       "peso_inicial": "Peso inicial (g)", "peso_final": "Peso final (g)", "metodos": "Métodos"},
            formatos={"importe_riesgo": MXN, "tasa": "0.000", "score": "0.00"},
            semaforo={"severidad": SEMAFORO_SEVERIDAD}, titulo="Hallazgos por línea de consumo",
            totales=["importe_riesgo"])
    if not ag.empty:
        libro.hoja_tabla("Patrones", ag, columnas=[
            "severidad", "regla", "producto", "dimension", "clave", "hospital", "valor_base", "valor_reciente", "ratio",
            "n_base", "n_reciente", "importe_riesgo", "motivo"],
            etiquetas={"clave": "Quién / dónde", "dimension": "Dimensión", "valor_base": "Valor base",
                       "valor_reciente": "Valor reciente", "n_base": "N base", "n_reciente": "N reciente",
                       "importe_riesgo": "Importe en riesgo"},
            formatos={"importe_riesgo": MXN, "ratio": "0.00", "valor_base": "0.000", "valor_reciente": "0.000"},
            semaforo={"severidad": SEMAFORO_SEVERIDAD}, titulo="Patrones agregados", totales=["importe_riesgo"])
    for dim, etiqueta in (("auxiliar", "Riesgo por auxiliar"), ("medico", "Riesgo por médico"),
                          ("hospital", "Riesgo por unidad"), ("subalmacen", "Riesgo por sub-almacén"),
                          ("producto", "Riesgo por producto")):
        filas = riesgos.get(dim) or []
        if filas:
            libro.hoja_tabla(etiqueta, pd.DataFrame(filas), columnas=[
                "clave", "indice", "hallazgos", "criticas", "altas", "importe_riesgo", "importe_total", "pct_riesgo", "lineas"],
                etiquetas={"clave": etiqueta.split(" por ")[1].capitalize(), "indice": "Índice de riesgo (0-100)",
                           "criticas": "Críticos", "altas": "Altos", "importe_riesgo": "Importe en riesgo",
                           "importe_total": "Importe consumido", "pct_riesgo": "% en riesgo", "lineas": "Líneas"},
                formatos={"importe_riesgo": MXN, "importe_total": MXN, "indice": "0.0", "pct_riesgo": "0.00"},
                titulo=etiqueta)
    if hist.get("tendencias_producto"):
        n = libro.hoja_tabla("Tendencias", pd.DataFrame(hist["tendencias_producto"]),
                             columnas=["producto", "meses", "tendencia_pct_mes", "prom_3m_previos", "prom_ult_3m",
                                       "variacion_pct", "importe_ult_3m"],
                             etiquetas={"tendencia_pct_mes": "Tendencia % / mes", "prom_3m_previos": "Prom. 3m previos",
                                        "prom_ult_3m": "Prom. últimos 3m", "variacion_pct": "Variación %",
                                        "importe_ult_3m": "Importe últimos 3m"},
                             formatos={"importe_ult_3m": MXN, "tendencia_pct_mes": "0.00", "variacion_pct": "0.0"},
                             titulo="Tendencia mensual por producto")
    if hist.get("comparativo_hospital"):
        libro.hoja_tabla("Comparativo unidades", pd.DataFrame(hist["comparativo_hospital"]),
                         etiquetas={"importe_30d": "Importe 30d", "importe_30d_prev": "Importe 30d previos",
                                    "variacion_pct": "Variación %", "folios_30d": "Folios 30d",
                                    "folios_30d_prev": "Folios 30d previos", "importe_por_folio": "Importe/folio",
                                    "importe_por_folio_prev": "Importe/folio previo"},
                         formatos={"importe_30d": MXN, "importe_30d_prev": MXN, "importe_por_folio": MXN,
                                   "importe_por_folio_prev": MXN},
                         titulo="Últimos 30 días vs. 30 anteriores por unidad médica")
    if hist.get("mensual"):
        m = pd.DataFrame(hist["mensual"])
        hoja = libro.hoja_tabla("Serie mensual", m, formatos={"importe": MXN, "ml_anestesicos_volatiles": DEC},
                                etiquetas={"ml_anestesicos_volatiles": "mL anestésicos volátiles", "lineas": "Líneas"},
                                titulo="Consumo mensual: importe (todas las unidades) y mL de anestésicos volátiles", como_tabla=False)
        if len(m) >= 3:
            libro.grafica(hoja, "linea", "Importe mensual consumido", 1, [2], 5, 4 + len(m), ancla="G4", eje_y="MXN")
    if hist.get("estacionalidad_semana"):
        e = pd.DataFrame([{"dia": k, "indice": v} for k, v in hist["estacionalidad_semana"].items()])
        libro.hoja_tabla("Estacionalidad", e, etiquetas={"dia": "Día", "indice": "Índice (1 = promedio)"},
                         formatos={"indice": "0.00"}, titulo="Índice de consumo por día de la semana")
    ruta = _ruta("anomalias", titulo)
    return _registrar(libro, ruta, "anomalias", titulo, {"corrida": (corrida or {}).get("id")}, int(len(h)), usuario)


# ── Agente 2 ────────────────────────────────────────────────────────────────
def reporte_pronostico(resultado: dict, informe: str = "", corrida: dict | None = None,
                       usuario: str = "") -> dict:
    r: pd.DataFrame = resultado.get("resurtido", pd.DataFrame())
    p: pd.DataFrame = resultado.get("pronosticos", pd.DataFrame())
    abc: pd.DataFrame = resultado.get("abc_xyz", pd.DataFrame())
    cad: pd.DataFrame = resultado.get("caducidades", pd.DataFrame())
    reb: pd.DataFrame = resultado.get("rebalanceo", pd.DataFrame())
    res = resultado.get("resumen", {})
    hist = resultado.get("historico", {})
    titulo = "Pronóstico de Demanda y Resurtido"
    libro = LibroExcel(titulo, f"Agente 2 · Corrida #{(corrida or {}).get('id', '')} · horizonte "
                               f"{res.get('horizonte_dias', '')} días · nivel {resultado.get('dimension', '')}")
    crit = res.get("criticidad", {})
    libro.portada(
        kpis=[("Combinaciones producto × ubicación", int(res.get("combinaciones", 0))),
              ("En desabasto", int(crit.get("desabasto", 0))), ("Críticas (< lead time)", int(crit.get("critico", 0))),
              ("Por reordenar", int(crit.get("reordenar", 0))), ("En exceso", int(crit.get("exceso", 0))),
              ("Compra sugerida a proveedor (MXN)", float(res.get("importe_compra_sugerida", 0))),
              ("Resurtido interno sugerido (MXN)", float(res.get("importe_resurtido_interno", 0))),
              ("Valor de inventario en red (MXN)", float(res.get("valor_inventario_total", 0))),
              ("Capital inmovilizado en exceso (MXN)", float(res.get("valor_inventario_exceso", 0))),
              ("Caducidades en riesgo (MXN)", float(res.get("importe_caducidad_en_riesgo", 0))),
              ("Transferencias propuestas", int(res.get("transferencias_propuestas", 0))),
              ("Error del pronóstico (WAPE semanal ponderado %)", float(res.get("wape_ponderado") or res.get("wape_promedio") or 0)),
              ("Unidades en camino consideradas", float(res.get("en_camino_unidades", 0))),
              ("Entregas retrasadas", int(res.get("entregas_retrasadas", 0))),
              ("Demanda de la agenda (folios programados, en su fecha)", float(res.get("demanda_comprometida_total", 0))),
              ("Ajuste atribuible a la agenda sobre el pronóstico", float(res.get("ajuste_agenda_total", 0))),
              ("Combinaciones con histórico insuficiente", int(res.get("historico_insuficiente", 0))),
              ("Importe PROPUESTO pendiente de decisión (MXN)", float(((resultado.get("importes") or {}).get("propuesto") or 0))),
              ("Importe APROBADO (MXN)", float(((resultado.get("importes") or {}).get("aprobado") or 0))),
              ("Importe EJECUTADO en Odoo (MXN)", float(((resultado.get("importes") or {}).get("ejecutado") or 0)))],
        notas=["Cada combinación producto × ubicación se pronostica con el método que mejor resultó en backtesting "
               "(origen móvil) y se calcula stock de seguridad al nivel de servicio configurado.",
               "La agenda (folios programados) es un piso diario del pronóstico en la fecha real de cada procedimiento; "
               "el saldo se proyecta al cierre de cada día con entradas y salidas fechadas; la cobertura son días hasta agotar.",
               "Sugerido local = lo necesario para no bajar del stock de seguridad dentro del ciclo de revisión + lead time interno. "
               "Sugerido de red = ídem con lead time del proveedor, en la unidad de compra (enteros).",
               "Las transferencias respetan la asignación conjunta: descuentan lo ya comprometido por otras propuestas y la reserva "
               "operativa del origen. Sugerido (motor) ≠ propuesto ≠ aprobado ≠ ejecutado: se informan por separado.",
               "Las cantidades se expresan en la unidad base del producto con su precisión (mL con decimal; piezas, pares, frascos y cajas enteros)."])
    if informe:
        libro.hoja_texto("Plan", informe, "Plan de abastecimiento del Agente 2")
    if not r.empty:
        libro.hoja_tabla("Resurtido", r[r["nivel"] == "local"], columnas=[
            "criticidad", "confianza", "producto", "almacen", "hospital", "unidad", "stock_actual", "reservado", "en_camino", "por_salir",
            "stock_proyectado", "demanda_diaria", "demanda_comprometida", "ajuste_agenda", "dias_con_agenda", "dias_cobertura", "fecha_quiebre",
            "fecha_necesaria", "lead_time_dias", "stock_seguridad",
            "punto_reorden", "demanda_horizonte", "sugerido", "importe_sugerido", "valor_inventario", "ultimos_28d", "prev_28d",
            "metodo", "mape", "sesgo_pct", "dias_imputados", "intermitente"],
            etiquetas={"almacen": "Ubicación", "stock_actual": "Existencia utilizable", "reservado": "Reservado", "en_camino": "En camino",
                       "por_salir": "Por salir", "stock_proyectado": "Existencia proyectada", "demanda_diaria": "Demanda/día",
                       "demanda_comprometida": "Demanda de la agenda", "ajuste_agenda": "Ajuste por agenda", "dias_con_agenda": "Días con agenda",
                       "dias_cobertura": "Cobertura (días hasta agotar)", "fecha_quiebre": "Fecha de quiebre", "fecha_necesaria": "Fecha necesaria de abasto",
                       "lead_time_dias": "Lead time", "stock_seguridad": "Stock seguridad", "punto_reorden": "Punto de reorden",
                       "demanda_horizonte": "Demanda horizonte", "importe_sugerido": "Importe sugerido", "valor_inventario": "Valor inventario",
                       "ultimos_28d": "Últimos 28d", "prev_28d": "28d previos", "metodo": "Método", "mape": "Error WAPE %",
                       "sesgo_pct": "Sesgo %", "dias_imputados": "Días imputados"},
            formatos={"importe_sugerido": MXN, "valor_inventario": MXN, "demanda_diaria": "0.000", "mape": "0.0"},
            semaforo={"criticidad": SEMAFORO_CRITICIDAD}, titulo="Resurtido por ubicación",
            totales=["importe_sugerido", "valor_inventario"])
        libro.hoja_tabla("Compras (red)", r[r["nivel"] == "red"], columnas=[
            "criticidad", "confianza", "producto", "unidad", "stock_actual", "reservado", "en_camino", "stock_proyectado", "demanda_diaria",
            "demanda_comprometida", "dias_cobertura", "fecha_quiebre", "fecha_necesaria", "lead_time_dias", "fecha_llegada_estimada", "llega_a_tiempo",
            "stock_seguridad", "punto_reorden", "demanda_horizonte",
            "sugerido", "sugerido_compra", "unidad_compra", "conversion_faltante", "costo_unit", "importe_sugerido", "valor_inventario", "mape"],
            etiquetas={"stock_actual": "Existencia utilizable red", "reservado": "Reservado", "en_camino": "Compras en camino",
                       "stock_proyectado": "Existencia proyectada", "demanda_diaria": "Demanda/día red",
                       "demanda_comprometida": "Demanda de la agenda", "dias_cobertura": "Cobertura (días hasta agotar)",
                       "fecha_quiebre": "Fecha de quiebre", "fecha_necesaria": "Fecha necesaria", "fecha_llegada_estimada": "Llegaría el",
                       "llega_a_tiempo": "¿Llega antes del quiebre?", "conversion_faltante": "Falta conversión de unidad",
                       "lead_time_dias": "Lead time proveedor", "stock_seguridad": "Stock seguridad", "punto_reorden": "Punto de reorden",
                       "demanda_horizonte": "Demanda horizonte", "sugerido": "Compra sugerida (unidad base)",
                       "sugerido_compra": "Compra sugerida (unidad de compra)", "unidad_compra": "Unidad de compra",
                       "costo_unit": "Costo unit.", "importe_sugerido": "Importe compra", "valor_inventario": "Valor inventario",
                       "mape": "Error WAPE %"},
            formatos={"importe_sugerido": MXN, "valor_inventario": MXN, "costo_unit": MXN, "demanda_diaria": "0.000"},
            semaforo={"criticidad": SEMAFORO_CRITICIDAD}, titulo="Compra sugerida a proveedor (nivel red)",
            totales=["importe_sugerido", "valor_inventario"])
    retr = resultado.get("retrasadas", pd.DataFrame())
    if retr is not None and not retr.empty:
        libro.hoja_tabla("Entregas retrasadas", retr, etiquetas={"ref": "Documento", "ubicacion_destino": "Destino",
                         "ubicacion_origen": "Origen / proveedor", "fecha_prevista": "Fecha prevista"},
                         titulo="Compras y transferencias con fecha prevista vencida (no se cuentan como abastecimiento)")
    if not reb.empty:
        libro.hoja_tabla("Transferencias", reb, columnas=[
            "criticidad_destino", "producto", "origen", "destino", "cantidad", "unidad", "cobertura_destino_dias",
            "cobertura_resultante_dias", "fecha_quiebre_destino", "fecha_necesaria", "fecha_llegada_estimada", "llega_a_tiempo",
            "stock_origen", "reserva_origen", "comprometido_origen_previo", "disponible_origen", "cubierto_previo_destino",
            "cobertura_origen_resultante", "importe", "motivo"],
            etiquetas={"criticidad_destino": "Criticidad destino", "cobertura_destino_dias": "Cobertura destino sin intervención (días)",
                       "cobertura_resultante_dias": "Cobertura destino con la transferencia (días)", "stock_origen": "Existencia origen",
                       "reserva_origen": "Reserva operativa del origen", "comprometido_origen_previo": "Ya comprometido en otras propuestas",
                       "disponible_origen": "Origen puede ceder", "cubierto_previo_destino": "Ya propuesto al destino",
                       "fecha_quiebre_destino": "Quiebre destino", "fecha_necesaria": "Fecha necesaria", "fecha_llegada_estimada": "Llegaría el",
                       "llega_a_tiempo": "¿Llega a tiempo?", "cobertura_origen_resultante": "Cobertura origen después (días)"},
            formatos={"importe": MXN}, semaforo={"criticidad_destino": SEMAFORO_CRITICIDAD},
            titulo="Transferencias internas propuestas", totales=["importe"])
    if not cad.empty:
        libro.hoja_tabla("Caducidades", cad, etiquetas={
            "almacen": "Ubicación", "dias_para_caducar": "Días para caducar", "demanda_diaria": "Demanda/día",
            "en_riesgo": "Unidades en riesgo", "importe_en_riesgo": "Importe en riesgo"},
            formatos={"importe_en_riesgo": MXN, "demanda_diaria": "0.000"},
            semaforo={"estado": SEMAFORO_ESTADO_CAD}, titulo="Lotes en riesgo de caducar (FEFO)",
            totales=["importe_en_riesgo"])
    if not abc.empty:
        libro.hoja_tabla("ABC-XYZ", abc, columnas=["clase", "abc", "xyz", "producto", "importe", "pct", "acum", "cv_semanal", "politica"],
                         etiquetas={"importe": "Importe 12m", "pct": "% del total", "acum": "% acumulado",
                                    "cv_semanal": "CV semanal", "politica": "Política sugerida"},
                         formatos={"importe": MXN, "pct": PCT, "acum": PCT, "cv_semanal": "0.00"},
                         titulo="Clasificación ABC (valor) × XYZ (variabilidad)")
    if not p.empty:
        piv = p[p["almacen"] != "RED (todas las ubicaciones)"].pivot_table(index="fecha", columns="producto",
                                                                             values="pronostico", aggfunc="sum").reset_index()
        hoja = libro.hoja_tabla("Pronóstico diario", piv, titulo="Pronóstico diario por producto (suma de ubicaciones)",
                                como_tabla=False, formatos={c: DEC for c in piv.columns if c != "fecha"})
        if len(piv) >= 5 and piv.shape[1] > 1:
            libro.grafica(hoja, "linea", "Pronóstico diario", 1, list(range(2, min(piv.shape[1], 7) + 1)), 5, 4 + len(piv),
                          ancla="B%d" % (8 + len(piv)), ancho=28, alto=12)
        libro.hoja_tabla("Pronóstico detalle", p, formatos={"pronostico": "0.000", "inferior": "0.000", "superior": "0.000"},
                         etiquetas={"almacen": "Ubicación", "mape": "WAPE %"}, titulo="Pronóstico por producto y ubicación")
    if hist.get("mensual"):
        m = pd.DataFrame(hist["mensual"])
        hoja = libro.hoja_tabla("Histórico mensual", m, formatos={"importe": MXN, "cantidad": DEC},
                                titulo="Demanda histórica mensual", como_tabla=False)
        if len(m) >= 3:
            libro.grafica(hoja, "barras", "Cantidad consumida por mes", 1, [2], 5, 4 + len(m), ancla="E4")
    if hist.get("crecimiento_90d"):
        libro.hoja_tabla("Crecimiento 90d", pd.DataFrame(hist["crecimiento_90d"]),
                         etiquetas={"ult_90d": "Últimos 90d", "prev_90d": "90d previos", "variacion_pct": "Variación %"},
                         formatos={"variacion_pct": "0.0"}, titulo="Productos que más cambian (90 días vs. 90 previos)")
    ruta = _ruta("pronostico", titulo)
    return _registrar(libro, ruta, "pronostico", titulo, {"corrida": (corrida or {}).get("id")}, int(len(r)), usuario)


# ── libre (copiloto) ────────────────────────────────────────────────────────
def reporte_libre(titulo: str, hojas: list[dict], notas: list[str] | None = None, kpis: list | None = None,
                  usuario: str = "") -> dict:
    """hojas = [{"nombre": str, "columnas": [...], "filas": [[...], ...] | [{...}, ...], "formatos": {...}}]"""
    libro = LibroExcel(titulo, "Generado por el Copiloto a solicitud del usuario")
    libro.portada(kpis=[(k, v) for k, v in (kpis or [])] or None, notas=notas or [])
    total = 0
    for h in hojas:
        filas = h.get("filas") or []
        cols = h.get("columnas")
        if filas and isinstance(filas[0], dict):
            df = pd.DataFrame(filas)
            if cols:
                df = df[[c for c in cols if c in df.columns]]
        else:
            df = pd.DataFrame(filas, columns=cols)
        total += len(df)
        libro.hoja_tabla(h.get("nombre", "Datos"), df, formatos=h.get("formatos") or {},
                         titulo=h.get("titulo") or h.get("nombre"), totales=h.get("totales"))
        if h.get("texto"):
            libro.hoja_texto(f"{h.get('nombre', 'Datos')} · notas", h["texto"])
    ruta = _ruta("libre", titulo)
    return _registrar(libro, ruta, "libre", titulo, {"hojas": [h.get("nombre") for h in hojas]}, total, usuario)


def reporte_dataframe_grande(titulo: str, df: pd.DataFrame, tipo: str = "consulta", hoja: str = "Datos",
                             metadatos: dict | None = None, usuario: str = "", limite: int = 1_000_000) -> dict:
    """Ruta rápida para reportes grandes (> 50,000 filas): libro en modo escritura secuencial, hasta el límite de Excel."""
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font, PatternFill
    from .excel import MORADO, BLANCO
    recortado = len(df) > limite
    d = df.head(limite)
    wb = Workbook(write_only=True)
    ws0 = wb.create_sheet("Acerca de")
    meta = {"Título": titulo, "Generado": datetime.now().strftime("%d-%b-%Y %H:%M"), "Registros": f"{len(d):,}",
            "Completo": ("NO — recortado al límite de Excel" if recortado else "sí"), "Moneda": "MXN"}
    meta.update(metadatos or {})
    for k, v in meta.items():
        ws0.append([k, str(v)])
    ws = wb.create_sheet(nombre_hoja_seguro(hoja))
    enc = []
    for c in d.columns:
        cell = WriteOnlyCell(ws, value=str(c).replace("_", " ").capitalize())
        cell.font = Font(bold=True, color=BLANCO); cell.fill = PatternFill("solid", start_color=MORADO, end_color=MORADO)
        enc.append(cell)
    ws.append(enc)
    for fila in d.itertuples(index=False):
        ws.append([_valor_plano(v) for v in fila])
    ruta = _ruta(tipo, titulo)
    ruta.parent.mkdir(parents=True, exist_ok=True)
    wb.save(ruta)
    rid = db.registrar_reporte(ruta.name, str(ruta), tipo, titulo, meta, int(len(d)), ruta.stat().st_size, usuario)
    db.log("info", "reportes", f"Excel grande generado: {ruta.name}", f"filas={len(d):,}", usuario)
    return {"id": rid, "archivo": ruta.name, "ruta": str(ruta), "tipo": tipo, "titulo": titulo, "filas": int(len(d)),
            "url": f"/reportes/{rid}/descargar", "recortado": recortado}


def nombre_hoja_seguro(s: str) -> str:
    return re.sub(r"[\[\]\*\?/\\:]", " ", s).strip()[:31] or "Datos"


def _valor_plano(v):
    if v is None:
        return None
    try:
        if pd.isna(v):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(v, pd.Timestamp):
        return v.to_pydatetime().replace(tzinfo=None)
    if hasattr(v, "item"):
        try:
            return v.item()
        except (ValueError, AttributeError):
            pass
    if isinstance(v, (list, dict, tuple, set)):
        return str(v)
    return v


def reporte_dataframe(titulo: str, df: pd.DataFrame, tipo: str = "consulta", hoja: str = "Datos",
                      formatos: dict | None = None, notas: list[str] | None = None, usuario: str = "",
                      semaforo: dict | None = None, metadatos: dict | None = None, total_origen: int | None = None) -> dict:
    if len(df) > 50_000:
        return reporte_dataframe_grande(titulo, df, tipo, hoja, metadatos, usuario)
    libro = LibroExcel(titulo)
    notas = list(notas or [])
    meta = dict(metadatos or {})
    if total_origen is not None and total_origen > len(df):
        notas.append(f"REPORTE INCOMPLETO: se muestran {len(df):,} de {total_origen:,} registros. Acota el periodo o los filtros.")
        meta["Completo"] = f"NO — {len(df):,} de {total_origen:,} registros"
    if "unidad" in df.columns:
        meta["Unidades presentes"] = ", ".join(sorted(set(df["unidad"].astype(str)) - {""}))
    libro.portada(kpis=[("Registros", int(len(df)))], notas=notas)
    libro.hoja_tabla(hoja, df, formatos=formatos or {}, titulo=titulo, semaforo=semaforo)
    ruta = _ruta(tipo, titulo)
    return _registrar(libro, ruta, tipo, titulo, meta.get("Parámetros", {}) if isinstance(meta.get("Parámetros"), dict) else {}, int(len(df)), usuario, meta)


def reporte_acciones(acciones: list[dict], usuario: str = "") -> dict:
    df = pd.DataFrame(acciones)
    if not df.empty:
        df["impacto"] = df["impacto"].map(lambda x: json.dumps(x, ensure_ascii=False) if isinstance(x, dict) else x)
        df["payload"] = df["payload"].map(lambda x: json.dumps(x, ensure_ascii=False) if isinstance(x, dict) else x)
    return reporte_dataframe("Cola de acciones autónomas", df, tipo="acciones", hoja="Acciones",
                             semaforo={"estado": SEMAFORO_ACCION}, usuario=usuario)
