"""Constructor de libros Excel con identidad visual de Odoo.

Cada libro lleva portada con KPIs, hojas tabulares con filtros, paneles
congelados, formatos numéricos, semáforos por columna y gráficas nativas.
"""
from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

# ── Paleta Odoo ─────────────────────────────────────────────────────────────
MORADO = "714B67"       # o-brand-odoo
MORADO_CLARO = "8F6B85"
TEAL = "017E84"         # o-brand-primary
TEAL_CLARO = "00A09D"
GRIS_OSCURO = "212529"
GRIS = "6C757D"
GRIS_CLARO = "F8F9FA"
BORDE = "DEE2E6"
VERDE = "28A745"
AMARILLO = "F0AD4E"
ROJO = "DC3545"
NARANJA = "E46E78"
BLANCO = "FFFFFF"

SEMAFORO_SEVERIDAD = {"critica": ROJO, "alta": NARANJA, "media": AMARILLO, "baja": TEAL_CLARO}
SEMAFORO_CRITICIDAD = {"desabasto": ROJO, "critico": NARANJA, "reordenar": AMARILLO, "ok": VERDE,
                       "exceso": MORADO_CLARO, "sin_movimiento": GRIS, "fuente": TEAL_CLARO,
                       "sin_existencias_registradas": GRIS}
SEMAFORO_ESTADO_CAD = {"caducado": ROJO, "critico": NARANJA, "riesgo": AMARILLO}
SEMAFORO_ACCION = {"propuesta": AMARILLO, "aprobada": TEAL_CLARO, "ejecutada": VERDE, "rechazada": GRIS,
                   "error": ROJO, "revertida": GRIS, "bloqueada": ROJO}

_fino = Side(style="thin", color=BORDE)
BORDE_FINO = Border(left=_fino, right=_fino, top=_fino, bottom=_fino)


def _fill(hex_: str) -> PatternFill:
    return PatternFill("solid", start_color=hex_, end_color=hex_)


def nombre_hoja(s: str) -> str:
    s = re.sub(r"[\[\]\*\?/\\:]", " ", s).strip()
    return (s or "Hoja")[:31]


class LibroExcel:
    MAX_FILAS = 1_000_000
    def __init__(self, titulo: str, subtitulo: str = "", autor: str = "CBH · Agentes de IA",
                 cliente: str = "Grupo CB · CBH+") -> None:
        self.recortes: list[dict] = []
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.titulo, self.subtitulo, self.autor, self.cliente = titulo, subtitulo, autor, cliente
        self.wb.properties.creator = autor
        self.wb.properties.title = titulo
        self._nombres: set[str] = set()

    # ── portada ─────────────────────────────────────────────────────────────
    def portada(self, kpis: list[tuple[str, Any]] | None = None, notas: Iterable[str] = (),
                secciones: list[tuple[str, list[tuple[str, Any]]]] | None = None) -> None:
        ws = self.wb.create_sheet(nombre_hoja("Portada"), 0)
        ws.sheet_view.showGridLines = False
        ws.column_dimensions["A"].width = 3
        ws.column_dimensions["B"].width = 44
        ws.column_dimensions["C"].width = 28
        ws.column_dimensions["D"].width = 28
        for r in range(1, 5):
            for c in range(1, 8):
                ws.cell(r, c).fill = _fill(MORADO)
        ws["B2"] = self.titulo
        ws["B2"].font = Font(size=20, bold=True, color=BLANCO)
        ws["B3"] = self.subtitulo or self.cliente
        ws["B3"].font = Font(size=11, color=BLANCO)
        ws["B6"] = "Generado"
        ws["C6"] = datetime.now().strftime("%d-%b-%Y %H:%M")
        ws["B7"] = "Cliente"
        ws["C7"] = self.cliente
        ws["B8"] = "Elaborado por"
        ws["C8"] = f"{self.autor} · Ingeniería Cóndor"
        for r in (6, 7, 8):
            ws.cell(r, 2).font = Font(bold=True, color=GRIS)
        fila = 10
        if kpis:
            ws.cell(fila, 2, "Indicadores clave").font = Font(size=13, bold=True, color=MORADO)
            fila += 1
            for etiqueta, valor in kpis:
                ws.cell(fila, 2, etiqueta).font = Font(color=GRIS_OSCURO)
                c = ws.cell(fila, 3, valor)
                c.font = Font(bold=True, size=12, color=TEAL)
                c.alignment = Alignment(horizontal="right")
                if isinstance(valor, float):
                    c.number_format = "#,##0.00"
                elif isinstance(valor, int):
                    c.number_format = "#,##0"
                fila += 1
            fila += 1
        for titulo_sec, pares in (secciones or []):
            ws.cell(fila, 2, titulo_sec).font = Font(size=13, bold=True, color=MORADO)
            fila += 1
            for etiqueta, valor in pares:
                ws.cell(fila, 2, etiqueta)
                c = ws.cell(fila, 3, valor)
                c.alignment = Alignment(horizontal="right")
                if isinstance(valor, float):
                    c.number_format = "#,##0.00"
                fila += 1
            fila += 1
        notas = list(notas)
        if notas:
            ws.cell(fila, 2, "Notas").font = Font(size=13, bold=True, color=MORADO)
            fila += 1
            for n in notas:
                ws.cell(fila, 2, f"• {n}").alignment = Alignment(wrap_text=True, vertical="top")
                ws.merge_cells(start_row=fila, start_column=2, end_row=fila, end_column=4)
                ws.row_dimensions[fila].height = max(15, 15 * (1 + len(n) // 90))
                fila += 1

    # ── tablas ──────────────────────────────────────────────────────────────
    def hoja_tabla(self, nombre: str, datos: pd.DataFrame | list[dict], columnas: list[str] | None = None,
                   etiquetas: dict[str, str] | None = None, formatos: dict[str, str] | None = None,
                   semaforo: dict[str, dict[str, str]] | None = None, titulo: str | None = None,
                   ancho_max: int = 60, como_tabla: bool = True, totales: list[str] | None = None) -> str:
        df = datos.copy() if isinstance(datos, pd.DataFrame) else pd.DataFrame(datos)
        if columnas:
            df = df[[c for c in columnas if c in df.columns]]
        etiquetas = etiquetas or {}
        formatos = formatos or {}
        semaforo = semaforo or {}
        # límite de Excel: si se recorta, se dice de forma explícita en la hoja y en "Acerca de" (nunca en silencio)
        total_original = len(df)
        recortado = total_original > self.MAX_FILAS
        if recortado:
            df = df.head(self.MAX_FILAS)
            self.recortes.append({"hoja": nombre, "mostradas": int(len(df)), "totales": int(total_original)})
        nombre = self._unico(nombre_hoja(nombre))
        ws = self.wb.create_sheet(nombre)
        ws.sheet_view.showGridLines = False
        fila0 = 1
        if titulo or recortado:
            ws.cell(1, 1, titulo or nombre).font = Font(size=14, bold=True, color=MORADO)
            aviso = f"{len(df):,} registros · {datetime.now().strftime('%d-%b-%Y %H:%M')}"
            if recortado:
                aviso = f"⚠ RECORTADO: se muestran {len(df):,} de {total_original:,} registros (límite de Excel). Pide el reporte con filtros o por partes."
            c2 = ws.cell(2, 1, aviso)
            c2.font = Font(color=("C00000" if recortado else GRIS), size=(11 if recortado else 9), bold=recortado)
            fila0 = 4
        cols = list(df.columns)
        for j, c in enumerate(cols, 1):
            celda = ws.cell(fila0, j, etiquetas.get(c, _bonito(c)))
            celda.font = Font(bold=True, color=BLANCO)
            celda.fill = _fill(MORADO)
            celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            celda.border = BORDE_FINO
        ws.row_dimensions[fila0].height = 30
        for i, fila in enumerate(df.itertuples(index=False), fila0 + 1):
            for j, v in enumerate(fila, 1):
                col = cols[j - 1]
                v = _valor_excel(v, str(col).lower())
                celda = ws.cell(i, j, v)
                celda.border = BORDE_FINO
                if isinstance(v, datetime):
                    celda.number_format = "yyyy-mm-dd hh:mm" if (v.hour or v.minute) else "yyyy-mm-dd"
                elif col in formatos:
                    celda.number_format = formatos[col]
                elif isinstance(v, float):
                    celda.number_format = "#,##0.00"
                elif isinstance(v, int) and not isinstance(v, bool):
                    celda.number_format = "#,##0"
                if col in semaforo and isinstance(v, str) and v in semaforo[col]:
                    celda.fill = _fill(semaforo[col][v])
                    celda.font = Font(bold=True, color=BLANCO if semaforo[col][v] not in (AMARILLO, TEAL_CLARO) else GRIS_OSCURO)
                    celda.alignment = Alignment(horizontal="center")
                elif i % 2 == 0:
                    celda.fill = _fill(GRIS_CLARO)
        ultima = fila0 + len(df)
        if totales and len(df):
            ws.cell(ultima + 1, 1, "TOTAL").font = Font(bold=True, color=MORADO)
            for j, c in enumerate(cols, 1):
                if c in totales:
                    try:
                        total = float(pd.to_numeric(df[c], errors="coerce").fillna(0).sum())
                    except (TypeError, ValueError):
                        continue
                    cel = ws.cell(ultima + 1, j, total)
                    cel.font = Font(bold=True, color=MORADO)
                    cel.number_format = formatos.get(c, "#,##0.00")
        if como_tabla and len(df) and len(cols):
            ref = f"A{fila0}:{get_column_letter(len(cols))}{max(ultima, fila0 + 1)}"
            t = Table(displayName=self._nombre_tabla(nombre), ref=ref)
            t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=False)
            ws.add_table(t)
        ws.freeze_panes = ws.cell(fila0 + 1, 1)
        for j, c in enumerate(cols, 1):
            largos = [len(str(v)) for v in df[c].head(500).tolist() if v is not None and v == v] if len(df) else [0]
            p90 = float(pd.Series(largos).quantile(0.9)) if largos else 10.0
            ancho = min(ancho_max, max(10, int(max(p90, len(str(etiquetas.get(c, _bonito(c)))) * 1.1)) + 2))
            ws.column_dimensions[get_column_letter(j)].width = ancho
        return nombre

    # ── metadatos del reporte ───────────────────────────────────────────────
    def acerca_de(self, datos: dict) -> str:
        """Hoja con periodo, filtros, unidades, moneda, origen, registros y si el reporte está completo."""
        ws = self.wb.create_sheet(self._unico("Acerca de"), 1)
        ws.sheet_view.showGridLines = False
        ws.column_dimensions["A"].width = 34
        ws.column_dimensions["B"].width = 80
        ws["A1"] = "Acerca de este reporte"
        ws["A1"].font = Font(size=14, bold=True, color=MORADO)
        fila = 3
        base = {"Generado": datetime.now().strftime("%d-%b-%Y %H:%M"), "Moneda": "MXN", "Elaborado por": f"{self.autor} · Ingeniería Cóndor",
                "Completo": ("NO — hojas recortadas al límite de Excel: " + "; ".join(f"{r['hoja']} {r['mostradas']:,} de {r['totales']:,}" for r in self.recortes)
                             if self.recortes else "sí (todas las filas)")}
        for k, v in {**base, **datos}.items():
            ws.cell(fila, 1, k).font = Font(bold=True, color=GRIS)
            c = ws.cell(fila, 2, v if not isinstance(v, (list, dict)) else str(v))
            c.alignment = Alignment(wrap_text=True, vertical="top")
            fila += 1
        return ws.title

    # ── texto (informes narrativos) ─────────────────────────────────────────
    def hoja_texto(self, nombre: str, texto: str, titulo: str | None = None) -> str:
        nombre = self._unico(nombre_hoja(nombre))
        ws = self.wb.create_sheet(nombre)
        ws.sheet_view.showGridLines = False
        ws.column_dimensions["A"].width = 2
        ws.column_dimensions["B"].width = 120
        fila = 1
        if titulo:
            ws.cell(fila, 2, titulo).font = Font(size=14, bold=True, color=MORADO)
            fila += 2
        for linea in (texto or "").splitlines():
            l = linea.rstrip()
            celda = ws.cell(fila, 2)
            if l.startswith("# "):
                celda.value, celda.font = l[2:], Font(size=14, bold=True, color=MORADO)
            elif l.startswith("## "):
                celda.value, celda.font = l[3:], Font(size=12, bold=True, color=TEAL)
            elif l.startswith("### "):
                celda.value, celda.font = l[4:], Font(size=11, bold=True, color=GRIS_OSCURO)
            elif l.startswith(("- ", "* ", "• ")):
                celda.value = "• " + _sin_md(l[2:])
            elif re.match(r"^\d+\. ", l):
                celda.value = _sin_md(l)
            elif l.startswith("|"):
                celda.value = " ".join(p.strip() for p in l.strip("|").split("|"))
                celda.font = Font(name="Consolas", size=9)
                if set(l.replace("|", "").strip()) <= {"-", ":", " "}:
                    celda.value = None
            else:
                celda.value = _sin_md(l)
            celda.alignment = Alignment(wrap_text=True, vertical="top")
            if celda.value:
                ws.row_dimensions[fila].height = max(15, 15 * (1 + len(str(celda.value)) // 110))
            fila += 1
        return nombre

    # ── gráficas ────────────────────────────────────────────────────────────
    def grafica(self, hoja: str, tipo: str, titulo: str, col_categorias: int, cols_series: list[int],
                fila_ini: int, fila_fin: int, ancla: str = "H4", ancho: float = 22, alto: float = 11,
                eje_y: str = "") -> None:
        ws = self.wb[hoja]
        ch = LineChart() if tipo == "linea" else BarChart()
        ch.title, ch.width, ch.height = titulo, ancho, alto
        ch.y_axis.title = eje_y
        ch.style = 10
        for c in cols_series:
            ref = Reference(ws, min_col=c, min_row=fila_ini - 1, max_row=fila_fin)
            ch.add_data(ref, titles_from_data=True)
        cats = Reference(ws, min_col=col_categorias, min_row=fila_ini, max_row=fila_fin)
        ch.set_categories(cats)
        ws.add_chart(ch, ancla)

    # ── guardar ─────────────────────────────────────────────────────────────
    def guardar(self, ruta: str | Path) -> Path:
        ruta = Path(ruta)
        ruta.parent.mkdir(parents=True, exist_ok=True)
        self.wb.save(ruta)
        return ruta

    # ── internos ────────────────────────────────────────────────────────────
    def _unico(self, n: str) -> str:
        base, k = n, 2
        while n in self._nombres or n in self.wb.sheetnames:
            n = f"{base[:28]} {k}"
            k += 1
        self._nombres.add(n)
        return n

    def _nombre_tabla(self, hoja: str) -> str:
        t = "T_" + re.sub(r"[^A-Za-z0-9]", "", hoja)[:24]
        k = 1
        base = t
        existentes = {tb.displayName for ws in self.wb.worksheets for tb in ws.tables.values()}
        while t in existentes:
            t = f"{base}{k}"
            k += 1
        return t


def _bonito(col: str) -> str:
    return col.replace("_", " ").strip().capitalize()


def _sin_md(s: str) -> str:
    return re.sub(r"\*\*(.+?)\*\*", r"\1", re.sub(r"`(.+?)`", r"\1", s))


_RE_ISO = re.compile(r"^\d{4}-\d{2}-\d{2}([ T]\d{2}:\d{2}(:\d{2})?)?")


def _valor_excel(v: Any, col: str = "") -> Any:
    if v is None:
        return None
    if isinstance(v, str) and (any(k in col for k in ("fecha", "caducidad", "_en", "dia", "semana", "mes")) and _RE_ISO.match(v)):
        try:
            return datetime.fromisoformat(v[:19].replace(" ", "T").replace("+00:00", ""))
        except ValueError:
            return v
    try:
        if pd.isna(v):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(v, (pd.Timestamp, datetime)):
        return v.to_pydatetime().replace(tzinfo=None) if isinstance(v, pd.Timestamp) else v.replace(tzinfo=None)
    if isinstance(v, (list, dict, tuple, set)):
        return ", ".join(str(x) for x in v) if isinstance(v, (list, tuple, set)) else str(v)
    if hasattr(v, "item"):
        try:
            return v.item()
        except (ValueError, AttributeError):
            pass
    return v
